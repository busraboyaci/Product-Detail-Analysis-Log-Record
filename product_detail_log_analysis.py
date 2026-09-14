#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
product_detail_log_analysis.py
------------------------------
Reads a raw product log Excel file, groups the rows by barcode and reports,
for every barcode: total log count, first log date + status, last log date + status.

Raw file columns (by position):
    A -> barcode
    B -> log status
    C -> date

Output file columns:
    A -> Barcode
    B -> Total Log Count
    C -> First Log Date
    D -> First Log Status
    E -> Last Log Date
    F -> Last Log Status

All directories are read from a .env file, so they can be changed without
touching the code.

Usage (Python console / terminal):
    python product_detail_log_analysis.py
    python product_detail_log_analysis.py --input C:/logs/raw.xlsx
    python product_detail_log_analysis.py --input-dir C:/logs --output-dir C:/reports
"""

from __future__ import annotations

import argparse
import os
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------- #
# 1. .env loading
# --------------------------------------------------------------------------- #

ENV_FILE = Path(__file__).resolve().parent / ".env"


def load_env(env_path: Path = ENV_FILE) -> None:
    """Load key=value pairs from .env into os.environ.

    Uses python-dotenv when installed; otherwise falls back to a small
    built-in parser so the script has no hard third-party dependency.
    """
    if not env_path.exists():
        print(f"[WARN] .env file not found: {env_path}")
        print("[WARN] Falling back to system environment variables / defaults.")
        return

    try:
        from dotenv import load_dotenv  # type: ignore

        load_dotenv(env_path, override=False)
        return
    except ImportError:
        pass

    with env_path.open("r", encoding="utf-8-sig") as fh:
        for raw_line in fh:
            line = raw_line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            key = key.strip()
            value = value.strip().strip('"').strip("'")
            if key and key not in os.environ:
                os.environ[key] = value


def env_str(key: str, default: str = "") -> str:
    value = os.getenv(key)
    return default if value is None or value.strip() == "" else value.strip()


def env_bool(key: str, default: bool) -> bool:
    value = os.getenv(key)
    if value is None or value.strip() == "":
        return default
    return value.strip().lower() in {"1", "true", "yes", "y", "evet", "on"}


def env_int(key: str, default: int) -> int:
    try:
        return int(env_str(key, str(default)))
    except ValueError:
        return default


# --------------------------------------------------------------------------- #
# 2. Configuration
# --------------------------------------------------------------------------- #

class Config:
    """Runtime configuration resolved from .env, then overridden by CLI args."""

    def __init__(self, args: argparse.Namespace) -> None:
        self.input_dir = Path(args.input_dir or env_str("INPUT_DIR", "./input")).expanduser()
        self.output_dir = Path(args.output_dir or env_str("OUTPUT_DIR", "./output")).expanduser()
        self.input_file = args.input or env_str("INPUT_FILE")
        self.output_prefix = env_str("OUTPUT_FILE_PREFIX", "product_detail_log_analysis")
        self.sheet_name = env_str("SHEET_NAME")           # empty -> first sheet
        self.output_sheet_name = env_str("OUTPUT_SHEET_NAME", "Analysis")
        self.has_header = env_bool("INPUT_HAS_HEADER", True)
        self.date_format = env_str("DATE_FORMAT", "%d.%m.%Y %H:%M:%S")
        self.dayfirst = env_bool("DATE_DAYFIRST", True)
        self.skip_rows = env_int("SKIP_ROWS", 0)
        self.timestamp_output = env_bool("TIMESTAMP_OUTPUT", True)

    def resolve_input_path(self) -> Path:
        """Return the Excel file to analyse.

        If INPUT_FILE is set, it is used (absolute, or relative to INPUT_DIR).
        Otherwise the most recently modified .xlsx/.xls file in INPUT_DIR is used.
        """
        if self.input_file:
            candidate = Path(self.input_file).expanduser()
            if not candidate.is_absolute():
                candidate = self.input_dir / candidate
            if not candidate.exists():
                raise FileNotFoundError(f"Input file not found: {candidate}")
            return candidate

        if not self.input_dir.exists():
            raise FileNotFoundError(f"Input directory not found: {self.input_dir}")

        files = [
            p for p in self.input_dir.iterdir()
            if p.is_file()
            and p.suffix.lower() in {".xlsx", ".xlsm", ".xls"}
            and not p.name.startswith("~$")
        ]
        if not files:
            raise FileNotFoundError(f"No Excel file found in directory: {self.input_dir}")

        newest = max(files, key=lambda p: p.stat().st_mtime)
        if len(files) > 1:
            print(f"[INFO] {len(files)} Excel files found; newest one selected: {newest.name}")
        return newest

    def build_output_path(self, input_path: Path) -> Path:
        self.output_dir.mkdir(parents=True, exist_ok=True)
        stem = f"{self.output_prefix}_{input_path.stem}"
        if self.timestamp_output:
            stem = f"{stem}_{datetime.now():%Y%m%d_%H%M%S}"
        return self.output_dir / f"{stem}.xlsx"


# --------------------------------------------------------------------------- #
# 3. Reading and cleaning
# --------------------------------------------------------------------------- #

COL_BARCODE, COL_STATUS, COL_DATE = "barcode", "log_status", "log_date"

HEADERS = [
    "Barcode",
    "Total Log Count",
    "First Log Date",
    "First Log Status",
    "Last Log Date",
    "Last Log Status",
]


def read_raw(path: Path, cfg: Config) -> pd.DataFrame:
    """Read the first three columns (A, B, C) by position, not by name."""
    sheet = cfg.sheet_name if cfg.sheet_name else 0
    df = pd.read_excel(
        path,
        sheet_name=sheet,
        header=0 if cfg.has_header else None,
        skiprows=cfg.skip_rows,
        usecols=[0, 1, 2],
        dtype=object,
        engine="openpyxl" if path.suffix.lower() in {".xlsx", ".xlsm"} else None,
    )
    if df.shape[1] < 3:
        raise ValueError(
            f"The file must have at least 3 columns (A, B, C); found {df.shape[1]}."
        )
    df.columns = [COL_BARCODE, COL_STATUS, COL_DATE]
    return df


def clean(df: pd.DataFrame, cfg: Config) -> pd.DataFrame:
    """Normalise barcode/status text, parse dates, drop unusable rows."""
    total_rows = len(df)

    # Keep the original row order as a deterministic tie-breaker.
    df = df.reset_index(drop=True)
    df["_row_order"] = df.index

    df[COL_BARCODE] = df[COL_BARCODE].map(normalize_barcode)
    df[COL_STATUS] = (
        df[COL_STATUS].astype(str).str.strip().replace({"nan": "", "None": "", "NaT": ""})
    )
    try:
        df[COL_DATE] = pd.to_datetime(
            df[COL_DATE], errors="coerce", dayfirst=cfg.dayfirst, format="mixed"
        )
    except (TypeError, ValueError):
        # pandas < 2.0 has no format="mixed"
        df[COL_DATE] = pd.to_datetime(df[COL_DATE], errors="coerce", dayfirst=cfg.dayfirst)

    no_barcode = df[COL_BARCODE].eq("")
    bad_date = df[COL_DATE].isna()

    if no_barcode.any():
        print(f"[WARN] {int(no_barcode.sum())} row(s) skipped: barcode is empty.")
    if (bad_date & ~no_barcode).any():
        print(f"[WARN] {int((bad_date & ~no_barcode).sum())} row(s) skipped: date is empty/invalid.")

    df = df[~no_barcode & ~bad_date].copy()
    print(f"[INFO] {len(df)} valid row(s) of {total_rows} will be analysed.")
    if df.empty:
        raise ValueError("No analysable rows remained after cleaning.")
    return df


def normalize_barcode(value: object) -> str:
    """Turn a barcode cell into clean text.

    Excel stores numeric barcodes as floats, so 8690000000001 can arrive as
    8690000000001.0 — that trailing '.0' is stripped here so the same physical
    product is never split into two groups.
    """
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    if text.lower() in {"nan", "none", "nat"}:
        return ""
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text


# --------------------------------------------------------------------------- #
# 4. Analysis
# --------------------------------------------------------------------------- #

def analyse(df: pd.DataFrame, cfg: Config) -> pd.DataFrame:
    """One output row per barcode: count, first/last date and their statuses."""
    ordered = df.sort_values(
        by=[COL_BARCODE, COL_DATE, "_row_order"], kind="mergesort"
    )
    grouped = ordered.groupby(COL_BARCODE, sort=True, as_index=False)

    result = grouped.agg(
        total_log_count=(COL_DATE, "size"),
        first_log_date=(COL_DATE, "first"),
        first_log_status=(COL_STATUS, "first"),
        last_log_date=(COL_DATE, "last"),
        last_log_status=(COL_STATUS, "last"),
    )

    result.columns = HEADERS
    result["First Log Date"] = result["First Log Date"].dt.strftime(cfg.date_format)
    result["Last Log Date"] = result["Last Log Date"].dt.strftime(cfg.date_format)
    return result


# --------------------------------------------------------------------------- #
# 5. Writing the Excel output
# --------------------------------------------------------------------------- #

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write_output(result: pd.DataFrame, out_path: Path, cfg: Config) -> None:
    """Write the result table; the header row is written exactly once."""
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        result.to_excel(writer, index=False, sheet_name=cfg.output_sheet_name)
        ws = writer.sheets[cfg.output_sheet_name]

        for cell in ws[1]:
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(HEADERS)):
            for cell in row:
                cell.font = BODY_FONT
                cell.border = BORDER
                if cell.column == 2:
                    cell.alignment = Alignment(horizontal="center")
                elif cell.column in (3, 5):
                    cell.alignment = Alignment(horizontal="center")

        for idx, header in enumerate(HEADERS, start=1):
            column = get_column_letter(idx)
            longest = max(
                [len(str(header))]
                + [len(str(v)) for v in result.iloc[:, idx - 1].head(500)]
            )
            ws.column_dimensions[column].width = min(max(longest + 4, 14), 40)

        ws.row_dimensions[1].height = 24
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"


# --------------------------------------------------------------------------- #
# 6. Entry point
# --------------------------------------------------------------------------- #

def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Produces a first/last log report per barcode from a product log Excel file."
    )
    parser.add_argument("--input", help="Excel file to analyse (overrides INPUT_FILE)")
    parser.add_argument("--input-dir", help="Input directory (overrides INPUT_DIR)")
    parser.add_argument("--output-dir", help="Output directory (overrides OUTPUT_DIR)")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    load_env()
    cfg = Config(parse_args(argv))

    try:
        input_path = cfg.resolve_input_path()
        print(f"[INFO] Input file : {input_path}")

        raw = read_raw(input_path, cfg)
        cleaned = clean(raw, cfg)
        result = analyse(cleaned, cfg)

        output_path = cfg.build_output_path(input_path)
        write_output(result, output_path, cfg)

        print(f"[OK]   {len(result)} unique barcode(s) reported.")
        print(f"[OK]   Output file: {output_path}")
        return 0

    except FileNotFoundError as exc:
        print(f"[ERROR] {exc}")
    except ValueError as exc:
        print(f"[ERROR] {exc}")
    except PermissionError:
        print("[ERROR] File is in use. Close the Excel file and try again.")
    except Exception as exc:  # noqa: BLE001 - console tool: report, don't crash
        print(f"[ERROR] Unexpected error: {type(exc).__name__}: {exc}")
    return 1


if __name__ == "__main__":
    sys.exit(main())
